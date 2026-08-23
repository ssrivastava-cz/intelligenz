"""Integration tests for `GET /api/v1/generation/{generation_id}/export/excel`
— downloads an Excel workbook built entirely from persisted generation
history. This endpoint's only dependencies are `HistoryService` and
`ExcelExportService`; it never touches OpenAI or ChromaDB at all, so
unlike other integration tests in this suite, there is nothing to fake
or override there.
"""
import io

import openpyxl

from app.core.dependencies import get_history_service
from app.main import app
from app.models.cost import ActualUsage, EstimatedUsage, MoneyAmount, PricingSnapshot
from app.models.generated_test_case import GeneratedTestCase, GeneratedTestStep
from app.models.generation_history import GenerationHistoryDraft, GenerationMetadata, RetrievalSummary
from app.services.history_service import HistoryService

_EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _override_history_service(tmp_path) -> HistoryService:
    service = HistoryService(history_root=tmp_path / "history")
    app.dependency_overrides[get_history_service] = lambda: service
    return service


def _clear_history_service_override() -> None:
    app.dependency_overrides.pop(get_history_service, None)


def _make_generated_test_case(**overrides) -> GeneratedTestCase:
    base = {
        "requirement_id": "REQ-1",
        "test_case_id": "TC-1",
        "test_case_title": "Reschedule an appointment",
        "priority": "High",
        "test_suite": "Appointments",
        "preconditions": "User is logged in.",
        "steps": [GeneratedTestStep(step_no=1, action="Open appointment.", expected_result="Details are shown.")],
        "post_conditions": "Appointment is updated.",
        "test_type": "Functional",
        "test_classification": "Smoke",
        "work_item_type": "Test Case",
        "automation_status": "Not Automated",
        "state": "Design",
        "area_path": "",
        "assigned_to": "",
        "tags": ["Appointments", "Regression", "Smoke"],
    }
    base.update(overrides)
    return GeneratedTestCase(**base)


def _make_generation_draft(test_cases: list[GeneratedTestCase] | None = None, **overrides) -> GenerationHistoryDraft:
    base = {
        "feature": "Appointments",
        "redmine_ticket": "12345",
        "model": "gpt-5",
        "prompt_version": "1.0.0",
        "retrieval": RetrievalSummary(
            workflow_chunks=1, historical_test_cases=1, historical_issues=0, uploaded_documents=0
        ),
        "estimated_usage": EstimatedUsage(
            model="gpt-5", input_tokens=1000, estimated_input_cost=MoneyAmount(usd=0.00125, inr=0.11)
        ),
        "actual_usage": ActualUsage(
            model="gpt-5",
            prompt_tokens=1000,
            completion_tokens=200,
            total_tokens=1200,
            input_cost=MoneyAmount(usd=0.00125, inr=0.11),
            output_cost=MoneyAmount(usd=0.002, inr=0.17),
            total_cost=MoneyAmount(usd=0.00325, inr=0.28),
        ),
        "pricing": PricingSnapshot(
            model="gpt-5",
            input_price_per_million_tokens=1.25,
            output_price_per_million_tokens=10.00,
            usd_to_inr_exchange_rate=87.00,
        ),
        "generation_time_ms": 2300.0,
        "prompt": "Prompt Version: 1.0.0\n\n...",
        "response": '{"testCases": [{"testCaseId": "TC-1"}]}',
        "test_cases": test_cases if test_cases is not None else [_make_generated_test_case()],
        "metadata": GenerationMetadata(top_k=5, upload_session_id="sess-abc123", generated_test_cases=1),
    }
    base.update(overrides)
    return GenerationHistoryDraft(**base)


async def test_export_excel_returns_200_with_correct_content_type(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        saved = history_service.save_generation_history(_make_generation_draft())

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        assert response.status_code == 200
        assert response.headers["content-type"] == _EXCEL_MEDIA_TYPE
    finally:
        _clear_history_service_override()


async def test_export_excel_content_disposition_has_the_expected_filename(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        saved = history_service.save_generation_history(_make_generation_draft())

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        expected_filename = f"ADO_TestCases_{saved.generation_id}.xlsx"
        assert response.headers["content-disposition"] == f'attachment; filename="{expected_filename}"'
    finally:
        _clear_history_service_override()


async def test_export_excel_body_is_a_valid_workbook_with_expected_content(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        saved = history_service.save_generation_history(_make_generation_draft())

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ["Test Cases", "Generation Summary"]

        test_cases_sheet = workbook["Test Cases"]
        header = [cell.value for cell in test_cases_sheet[1]]
        assert header[:3] == ["ID", "Work Item Type", "Title"]
        [header_row, step_row] = list(test_cases_sheet.iter_rows(min_row=2, values_only=True))
        assert header_row[0] is None  # ID — left for Azure DevOps to assign.
        assert header_row[1] == "Test Case"
        assert header_row[2] == "Reschedule an appointment"
        assert step_row[4] == "Open appointment."  # Step Action, on its own row

        summary_sheet = workbook["Generation Summary"]
        summary_values = {row[0]: row[1] for row in summary_sheet.iter_rows(min_row=2, values_only=True)}
        assert summary_values["Generation ID"] == saved.generation_id
        assert summary_values["Feature"] == "Appointments"
        assert summary_values["Number of Test Cases"] == 1
    finally:
        _clear_history_service_override()


async def test_export_excel_handles_multiple_test_cases_and_steps(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        test_cases = [
            _make_generated_test_case(
                test_case_id="TC-1",
                test_case_title="Verify code search in coding tool",
                steps=[GeneratedTestStep(step_no=1, action="Step 1.", expected_result="Result 1.")],
            ),
            _make_generated_test_case(
                test_case_id="TC-2",
                test_case_title="Verify provider search in coding tool",
                steps=[
                    GeneratedTestStep(step_no=1, action="Step 1.", expected_result="Result 1."),
                    GeneratedTestStep(step_no=2, action="Step 2.", expected_result="Result 2."),
                ],
            ),
        ]
        saved = history_service.save_generation_history(_make_generation_draft(test_cases=test_cases))

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = workbook["Test Cases"]
        data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
        # TC-1: 1 header row + 1 step row. TC-2: 1 header row + 2 step rows.
        assert len(data_rows) == 5
        titles = {row[2] for row in data_rows if row[2]}
        assert titles == {"Verify code search in coding tool", "Verify provider search in coding tool"}
    finally:
        _clear_history_service_override()


async def test_export_excel_writes_automation_status_as_exactly_not_automated(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        saved = history_service.save_generation_history(_make_generation_draft())

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        [header_row] = [
            row for row in workbook["Test Cases"].iter_rows(min_row=2, values_only=True) if row[1] == "Test Case"
        ]
        assert header_row[10] == "Not Automated"
        assert header_row[10] not in {"No", "NOT AUTOMATED", "Not automated", "Automated", "Yes"}
    finally:
        _clear_history_service_override()


async def test_export_excel_preserves_smoke_and_regression_only_tags_exactly(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        test_cases = [
            _make_generated_test_case(
                test_case_id="TC-1", tags=["Contact and Sticket Log", "Regression", "Smoke"]
            ),
            _make_generated_test_case(test_case_id="TC-2", tags=["Contact and Sticket Log", "Regression"]),
        ]
        saved = history_service.save_generation_history(_make_generation_draft(test_cases=test_cases))

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        header_rows = [
            row for row in workbook["Test Cases"].iter_rows(min_row=2, values_only=True) if row[1] == "Test Case"
        ]
        tags_by_row = [row[9] for row in header_rows]
        assert "Contact and Sticket Log; Regression; Smoke" in tags_by_row
        assert "Contact and Sticket Log; Regression" in tags_by_row
        assert "Contact and Sticket Log; Smoke" not in tags_by_row
    finally:
        _clear_history_service_override()


async def test_export_excel_renders_empty_tags_without_error(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        saved = history_service.save_generation_history(
            _make_generation_draft(test_cases=[_make_generated_test_case(tags=[], steps=[])])
        )

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        assert response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        [data_row] = list(workbook["Test Cases"].iter_rows(min_row=2, values_only=True))
        assert data_row[9] is None  # Tags
    finally:
        _clear_history_service_override()


async def test_export_excel_404_for_unknown_generation(client, tmp_path):
    _override_history_service(tmp_path)

    try:
        response = await client.get("/api/v1/generation/2000-01-01T00-00-00/export/excel")

        assert response.status_code == 404
    finally:
        _clear_history_service_override()


async def test_export_excel_409_when_generation_has_no_test_cases(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        empty_metadata = GenerationMetadata(top_k=5, upload_session_id=None, generated_test_cases=0)
        saved = history_service.save_generation_history(
            _make_generation_draft(test_cases=[], metadata=empty_metadata)
        )

        response = await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        assert response.status_code == 409
        assert response.json()["detail"]
    finally:
        _clear_history_service_override()


async def test_export_excel_does_not_modify_the_stored_history(client, tmp_path):
    history_service = _override_history_service(tmp_path)

    try:
        saved = history_service.save_generation_history(_make_generation_draft())

        await client.get(f"/api/v1/generation/{saved.generation_id}/export/excel")

        reloaded = history_service.load_generation(saved.generation_id)
        assert reloaded == saved
    finally:
        _clear_history_service_override()
