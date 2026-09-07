"""Extracts text from XLSX files, one section per sheet."""
import io

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.exceptions import ValidationError
from app.models.document import Document
from app.models.parsed_document import DocumentSection, ParsedDocument, ParsedDocumentMetadata
from app.services.parsers.base import DocumentParser
from app.utils.text import title_from_filename


class ExcelParser(DocumentParser):
    parser_name = "ExcelParser"
    parser_version = "1.0"

    def parse(self, document: Document, content: bytes) -> ParsedDocument:
        try:
            workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
            sections = [
                DocumentSection(heading=sheet_name, content=_render_sheet(workbook[sheet_name]))
                for sheet_name in workbook.sheetnames
            ]
        except Exception as exc:
            raise ValidationError(f"Failed to parse XLSX '{document.filename}': {exc}") from exc

        return ParsedDocument(
            document_id=document.id,
            title=title_from_filename(document.filename),
            sections=sections,
            content="\n\n".join(section.content for section in sections if section.content),
            metadata=ParsedDocumentMetadata(
                feature=document.feature,
                document_type=document.document_type,
                document_source=document.source,
                source_filename=document.filename,
                page_number=None,  # sheets, not pages
                parser_name=self.parser_name,
                parser_version=self.parser_version,
                source_path=document.source_relative_path,
                source_folder=document.source_folder,
            ),
        )


def _render_sheet(sheet: Worksheet) -> str:
    lines = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if value is None else str(value) for value in row]
        if any(cell for cell in cells):
            lines.append(" | ".join(cells))
    return "\n".join(lines)
