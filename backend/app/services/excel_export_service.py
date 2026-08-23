"""Converts a persisted `GenerationHistoryEntry` into an ADO-compatible
Excel workbook. The persisted history is the only data source: this
service never regenerates AI output, never performs retrieval, never
calls OpenAI, and never touches ChromaDB — it knows nothing about
`RetrievalService`, `PromptBuilder`, `GenerationService`, or the OpenAI
SDK, only the already-persisted `GenerationHistoryEntry` model.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from app.core.exceptions import ConflictError
from app.models.generated_test_case import GeneratedTestCase, GeneratedTestStep
from app.models.generation_history import GenerationHistoryEntry

_TEST_CASES_SHEET_NAME = "Test Cases"
_METADATA_SHEET_NAME = "Generation Summary"

# Matches Azure DevOps' own Test Case Excel import layout exactly: one
# "header" row per test case (Work Item Type/Title/Area Path/Assigned
# To/State/Tags/Automation Status; the step columns blank) followed by
# one row per step (Test Step/Step Action/Step Expected only; every
# test-case-level column blank) — never every field repeated on every
# row. `ID` is always left blank so ADO assigns it on import; this
# application never fabricates one.
_TEST_CASE_COLUMNS = [
    "ID",
    "Work Item Type",
    "Title",
    "Test Step",
    "Step Action",
    "Step Expected",
    "Area Path",
    "Assigned To",
    "State",
    "Tags",
    "Automation Status",
]

_WRAP_TEXT_COLUMNS = {"Title", "Step Action", "Step Expected"}
_CENTER_COLUMNS = {"Test Step", "State", "Automation Status"}


class ExcelExportService:
    def build_workbook(self, entry: GenerationHistoryEntry) -> bytes:
        """Raises `ConflictError` if the generation has no parsed test
        cases — there is nothing meaningful to export.
        """
        if not entry.test_cases:
            raise ConflictError(f"Generation '{entry.generation_id}' has no parsed test cases to export.")

        workbook = Workbook()

        test_cases_sheet = workbook.active
        test_cases_sheet.title = _TEST_CASES_SHEET_NAME
        _build_test_cases_sheet(test_cases_sheet, entry)

        metadata_sheet = workbook.create_sheet(_METADATA_SHEET_NAME)
        _build_metadata_sheet(metadata_sheet, entry)

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


def _build_test_cases_sheet(sheet: Worksheet, entry: GenerationHistoryEntry) -> None:
    sheet.append(_TEST_CASE_COLUMNS)
    _bold_row(sheet, row_index=1)
    sheet.freeze_panes = "A2"

    for test_case in entry.test_cases:
        # One header row per test case (test-case-level fields only,
        # step columns blank) — a test case with no steps still gets
        # exactly this one row, so it's never silently dropped from the
        # export — followed by one row per step (step columns only,
        # every test-case-level field blank). Never both sets of fields
        # on the same row: that's what lets Azure DevOps import this
        # sheet as one test case with N steps, not N separate test cases.
        sheet.append(_test_case_header_row(test_case))
        for step in test_case.steps:
            sheet.append(_test_step_row(step))

    _apply_column_alignment(sheet)
    _autosize_columns(sheet)


def _test_case_header_row(test_case: GeneratedTestCase) -> list:
    tags = "; ".join(test_case.tags) if test_case.tags else ""
    return [
        "",  # ID — left blank for Azure DevOps to assign; never fabricated.
        test_case.work_item_type or "",
        test_case.test_case_title,
        "",
        "",
        "",
        test_case.area_path or "",
        test_case.assigned_to or "",
        test_case.state or "",
        tags,
        test_case.automation_status or "",
    ]


def _test_step_row(step: GeneratedTestStep) -> list:
    return ["", "", "", step.step_no, step.action, step.expected_result, "", "", "", "", ""]


def _build_metadata_sheet(sheet: Worksheet, entry: GenerationHistoryEntry) -> None:
    actual_usage = entry.actual_usage
    rows = [
        ("Generation ID", entry.generation_id),
        ("Timestamp", entry.timestamp.isoformat()),
        ("Feature", entry.feature),
        ("Redmine Ticket", entry.redmine_ticket),
        ("Prompt Version", entry.prompt_version),
        ("LLM Model", entry.model),
        ("Generation Time (ms)", entry.generation_time_ms),
        ("Input Tokens", actual_usage.prompt_tokens),
        ("Output Tokens", actual_usage.completion_tokens),
        ("Total Tokens", actual_usage.total_tokens),
        ("Input Cost (USD)", actual_usage.input_cost.usd),
        ("Output Cost (USD)", actual_usage.output_cost.usd),
        ("Total Cost (USD)", actual_usage.total_cost.usd),
        ("Input Cost (INR)", actual_usage.input_cost.inr),
        ("Output Cost (INR)", actual_usage.output_cost.inr),
        ("Total Cost (INR)", actual_usage.total_cost.inr),
        ("Exchange Rate", entry.pricing.usd_to_inr_exchange_rate),
        ("Number of Test Cases", len(entry.test_cases)),
        ("Total Test Steps", sum(len(test_case.steps) for test_case in entry.test_cases)),
        ("Upload Session", entry.metadata.upload_session_id or ""),
    ]

    sheet.append(["Field", "Value"])
    _bold_row(sheet, row_index=1)
    for row in rows:
        sheet.append(row)

    _autosize_columns(sheet)


def _bold_row(sheet: Worksheet, row_index: int) -> None:
    for cell in sheet[row_index]:
        cell.font = Font(bold=True)


def _apply_column_alignment(sheet: Worksheet) -> None:
    for column_index, column_name in enumerate(_TEST_CASE_COLUMNS, start=1):
        wrap_text = column_name in _WRAP_TEXT_COLUMNS
        center = column_name in _CENTER_COLUMNS
        if not wrap_text and not center:
            continue

        alignment = Alignment(wrap_text=wrap_text, horizontal="center" if center else None)
        for row in sheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
            for cell in row:
                cell.alignment = alignment


def _autosize_columns(sheet: Worksheet) -> None:
    for column_cells in sheet.columns:
        content_length = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=0)
        sheet.column_dimensions[column_cells[0].column_letter].width = content_length + 2
