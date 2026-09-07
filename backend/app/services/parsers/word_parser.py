"""Extracts text from DOCX files, sectioned by heading style.

Tables are extracted too, not just paragraphs: `docx.tables` and
`docx.paragraphs` are separate flat lists on a python-docx `Document`,
which loses where a table actually sits relative to the surrounding
headings/paragraphs — `_iter_block_items` walks the underlying XML body
in document order instead, so a table lands in the same
`DocumentSection` (and at the same position) it actually appears in.
"""
import io
from collections.abc import Iterator

from docx import Document as open_docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.exceptions import ValidationError
from app.models.document import Document
from app.models.parsed_document import DocumentSection, ParsedDocument, ParsedDocumentMetadata
from app.services.parsers.base import DocumentParser
from app.utils.text import is_blank_row, title_from_filename

_HEADING_STYLE_PREFIX = "Heading"


class WordParser(DocumentParser):
    parser_name = "WordParser"
    parser_version = "1.0"

    def parse(self, document: Document, content: bytes) -> ParsedDocument:
        try:
            docx = open_docx(io.BytesIO(content))
            sections = _split_into_sections(docx)
            core_title = (docx.core_properties.title or "").strip()
        except Exception as exc:
            raise ValidationError(f"Failed to parse DOCX '{document.filename}': {exc}") from exc

        return ParsedDocument(
            document_id=document.id,
            title=core_title or title_from_filename(document.filename),
            sections=sections,
            content="\n\n".join(section.content for section in sections if section.content),
            metadata=ParsedDocumentMetadata(
                feature=document.feature,
                document_type=document.document_type,
                document_source=document.source,
                source_filename=document.filename,
                page_number=None,  # DOCX has no fixed page count without rendering
                parser_name=self.parser_name,
                parser_version=self.parser_version,
                source_path=document.source_relative_path,
                source_folder=document.source_folder,
            ),
        )


def _iter_block_items(docx: DocxDocument) -> Iterator[Paragraph | Table]:
    """Yields every top-level paragraph and table in the order they
    actually appear in the document — the standard python-docx technique
    for this, since `docx.paragraphs`/`docx.tables` are separate,
    order-losing flat lists.
    """
    for child in docx.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, docx)
        elif child.tag == qn("w:tbl"):
            yield Table(child, docx)


def _split_into_sections(docx: DocxDocument) -> list[DocumentSection]:
    sections: list[DocumentSection] = []
    current_heading: str | None = None
    current_paragraphs: list[str] = []
    table_number = 0

    def flush() -> None:
        body = "\n".join(p for p in current_paragraphs if p.strip())
        if current_heading or body:
            sections.append(DocumentSection(heading=current_heading, content=body))

    for block in _iter_block_items(docx):
        if isinstance(block, Table):
            table_number += 1
            current_paragraphs.extend(_render_table(table_number, block))
            continue

        style_name = block.style.name if block.style else ""
        if style_name.startswith(_HEADING_STYLE_PREFIX):
            flush()
            current_heading = block.text.strip()
            current_paragraphs = []
        else:
            current_paragraphs.append(block.text)

    flush()

    if not sections:
        sections.append(DocumentSection(heading=None, content=""))

    return sections


def _render_table(table_number: int, table: Table) -> list[str]:
    """Renders one table as a list of self-contained "Header: Value"
    blocks, one per non-blank data row — never one giant concatenated
    string. Each block repeats the table's own "Table N" marker and
    every column's header inline (matching the same "key: value"
    convention `CsvParser` already uses for one CSV row), so a chunk
    boundary that happens to fall between two rows of a large table
    never leaves either fragment without knowing which table or which
    headers it belongs to.

    The first row is always treated as the header row (the normal DOCX
    convention). A completely blank row (see `is_blank_row`) is
    skipped, just like a blank CSV row. A cell beyond the header's
    column count — or a header cell that's itself blank — falls back to
    "Column N" rather than being silently dropped.
    """
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    rows = [row for row in rows if not is_blank_row(row)]
    if not rows:
        return []

    header, *data_rows = rows
    if not data_rows:
        return [f"Table {table_number}\n" + " | ".join(header)]

    blocks = []
    for row in data_rows:
        pairs = [_label(header, index) + f": {value}" for index, value in enumerate(row)]
        blocks.append(f"Table {table_number}\n" + "\n".join(pairs))
    return blocks


def _label(header: list[str], index: int) -> str:
    if index < len(header) and header[index]:
        return header[index]
    return f"Column {index + 1}"
