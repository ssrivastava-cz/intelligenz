import io

import pytest
from docx import Document as open_docx

from app.core.exceptions import ValidationError
from app.models.common import DocumentCategory, DocumentSource, DocumentType
from app.services.chunking_engine import ChunkingEngine
from app.services.parsers.word_parser import WordParser
from tests.unit.parsers.factories import make_document


def _build_docx(sections: list[tuple[str, str]]) -> bytes:
    docx = open_docx()
    for heading, body in sections:
        docx.add_heading(heading, level=1)
        docx.add_paragraph(body)
    buffer = io.BytesIO()
    docx.save(buffer)
    return buffer.getvalue()


def test_splits_sections_by_heading():
    content = _build_docx([("Overview", "This describes the workflow."), ("Steps", "Do this then that.")])
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, content)

    assert [s.heading for s in parsed.sections] == ["Overview", "Steps"]
    assert "This describes the workflow." in parsed.sections[0].content
    assert "Do this then that." in parsed.sections[1].content
    assert "This describes the workflow." in parsed.content
    assert parsed.metadata.page_number is None


def test_metadata_reflects_user_upload_source():
    content = _build_docx([("Notes", "Body")])
    document = make_document(
        filename="notes.docx",
        document_type=DocumentType.DOCX,
        feature="Analytics",
        source=DocumentSource.USER_UPLOAD,
    )

    parsed = WordParser().parse(document, content)

    assert parsed.metadata.document_source == DocumentSource.USER_UPLOAD
    assert parsed.metadata.feature == "Analytics"
    assert parsed.metadata.parser_name == "WordParser"
    assert parsed.metadata.parser_version == "1.0"


def test_uses_document_core_title_when_present():
    docx = open_docx()
    docx.core_properties.title = "Workflow Guide"
    docx.add_paragraph("Body")
    buffer = io.BytesIO()
    docx.save(buffer)

    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)
    parsed = WordParser().parse(document, buffer.getvalue())

    assert parsed.title == "Workflow Guide"


def test_falls_back_to_filename_when_no_title_or_headings():
    content = _build_docx([])
    document = make_document(filename="misc_notes.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, content)

    assert parsed.title == "Misc Notes"


def test_raises_validation_error_for_corrupt_docx():
    document = make_document(filename="broken.docx", document_type=DocumentType.DOCX)

    with pytest.raises(ValidationError):
        WordParser().parse(document, b"not a real docx")


def test_metadata_carries_the_portable_source_path_and_folder():
    content = _build_docx([("Notes", "Body")])
    document = make_document(
        filename="notes.docx",
        document_type=DocumentType.DOCX,
        source_relative_path="source_of_truth/Service Model 1/workflows/notes.docx",
        source_folder="Service Model 1",
    )

    parsed = WordParser().parse(document, content)

    assert parsed.metadata.source_path == "source_of_truth/Service Model 1/workflows/notes.docx"
    assert parsed.metadata.source_folder == "Service Model 1"


# --- table extraction ---


def _add_table(docx, rows: list[list[str]]):
    table = docx.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            table.cell(row_index, col_index).text = value
    return table


def test_paragraphs_and_headings_still_parse_correctly_alongside_a_table():
    docx = open_docx()
    docx.add_heading("Overview", level=1)
    docx.add_paragraph("This describes the workflow.")
    _add_table(docx, [["Data Type", "Owner"], ["Provider Data", "Shivasish"]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    assert len(parsed.sections) == 1
    assert parsed.sections[0].heading == "Overview"
    assert "This describes the workflow." in parsed.sections[0].content


def test_table_cells_are_extracted_with_header_value_pairing():
    docx = open_docx()
    docx.add_heading("Data Sources", level=1)
    _add_table(
        docx,
        [
            ["Data Type", "Primary Source", "Owner"],
            ["Provider Data", "CCD", "Shivasish"],
        ],
    )
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    content = parsed.sections[0].content
    assert "Data Type: Provider Data" in content
    assert "Primary Source: CCD" in content
    assert "Owner: Shivasish" in content
    # Never a flat concatenation with no structure.
    assert "Provider Data CCD Shivasish" not in content


def test_multiple_tables_are_extracted_in_order():
    docx = open_docx()
    docx.add_heading("Section One", level=1)
    _add_table(docx, [["Header A"], ["Value 1"]])
    docx.add_heading("Section Two", level=1)
    _add_table(docx, [["Header B"], ["Value 2"]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    assert [s.heading for s in parsed.sections] == ["Section One", "Section Two"]
    assert "Header A: Value 1" in parsed.sections[0].content
    assert "Header B: Value 2" in parsed.sections[1].content
    # Table numbering reflects document-wide order, not per-section order.
    assert "Table 1" in parsed.sections[0].content
    assert "Table 2" in parsed.sections[1].content


def test_table_row_order_is_preserved():
    docx = open_docx()
    _add_table(
        docx,
        [["Step"], ["First"], ["Second"], ["Third"]],
    )
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    content = parsed.sections[0].content
    assert content.index("Step: First") < content.index("Step: Second") < content.index("Step: Third")


def test_each_table_row_is_self_contained_with_its_own_table_and_header_labels():
    """So a chunk boundary falling between two rows of a large table
    never leaves either fragment without knowing which table or which
    headers it belongs to (no shared header line printed once at the
    top with bare positional values after it)."""
    docx = open_docx()
    _add_table(docx, [["Scenario", "Action"], ["New Appt", "Create visit"], ["Cancelled Appt", "Log only"]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())
    content = parsed.sections[0].content

    assert content.count("Table 1") == 2
    assert "Scenario: New Appt\nAction: Create visit" in content
    assert "Scenario: Cancelled Appt\nAction: Log only" in content


def test_blank_table_rows_are_skipped():
    docx = open_docx()
    _add_table(docx, [["Header"], ["Value"], ["   "], [""]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())
    content = parsed.sections[0].content

    assert content.count("Header:") == 1


def test_a_table_with_only_a_header_row_renders_the_header_as_plain_text():
    docx = open_docx()
    _add_table(docx, [["Column A", "Column B"]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    assert "Column A | Column B" in parsed.sections[0].content


def test_a_completely_blank_table_produces_no_table_content():
    docx = open_docx()
    docx.add_paragraph("Body text.")
    _add_table(docx, [["", ""], ["", ""]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    assert "Table" not in parsed.sections[0].content
    assert "Body text." in parsed.sections[0].content


def test_table_content_reaches_chunking():
    docx = open_docx()
    docx.add_heading("Data Sources", level=1)
    _add_table(docx, [["Data Type", "Owner"], ["Provider Data", "Shivasish"]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())
    chunks = ChunkingEngine().chunk_document(parsed, DocumentCategory.WORKFLOW)

    assert len(chunks) == 1
    assert "Data Type: Provider Data" in chunks[0].chunk_text
    assert "Owner: Shivasish" in chunks[0].chunk_text


def test_table_content_flows_into_the_documents_full_text():
    docx = open_docx()
    _add_table(docx, [["Field"], ["Contact Log"]])
    buffer = io.BytesIO()
    docx.save(buffer)
    document = make_document(filename="workflow.docx", document_type=DocumentType.DOCX)

    parsed = WordParser().parse(document, buffer.getvalue())

    assert "Field: Contact Log" in parsed.content
