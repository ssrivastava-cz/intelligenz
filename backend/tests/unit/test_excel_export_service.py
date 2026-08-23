"""Unit tests for ExcelExportService — converts a persisted
`GenerationHistoryEntry` into an ADO-compatible Excel workbook. The
persisted history is the only input; there is no OpenAI, retrieval, or
ChromaDB dependency to fake here at all.

The "Test Cases" sheet matches Azure DevOps' own Test Case Excel import
layout: one header row per test case (Work Item Type/Title/Area Path/
Assigned To/State/Tags/Automation Status; step columns blank) followed
by one row per step (Test Step/Step Action/Step Expected; every
test-case-level column blank) — this is what lets ADO import the sheet
as one test case with N steps, not N separate test cases.
"""
import io
from datetime import UTC, datetime

import openpyxl
import pytest

from app.core.exceptions import ConflictError
from app.models.cost import ActualUsage, EstimatedUsage, MoneyAmount, PricingSnapshot
from app.models.generated_test_case import GeneratedTestCase, GeneratedTestStep
from app.models.generation_history import GenerationHistoryEntry, GenerationMetadata, RetrievalSummary
from app.services.excel_export_service import ExcelExportService

_TEST_CASE_COLUMNS = [
    "ID",
    "Work Item Type",
    "Title",
    "Test Step",
    "Step Action",
    "Step Expected",
    "Area Path",
    "Assigned To",
    "State",
    "Tags",
    "Automation Status",
]


def _make_step(**overrides) -> GeneratedTestStep:
    base = {"step_no": 1, "action": "Open appointment.", "expected_result": "Details are shown."}
    base.update(overrides)
    return GeneratedTestStep(**base)


def _make_test_case(**overrides) -> GeneratedTestCase:
    base = {
        "requirement_id": "REQ-1",
        "test_case_id": "TC-1",
        "test_case_title": "Reschedule an appointment",
        "priority": "High",
        "test_suite": "Appointments",
        "preconditions": "User is logged in.",
        "steps": [_make_step()],
        "post_conditions": "Appointment is updated.",
        "test_type": "Functional",
        "test_classification": "Smoke",
        "work_item_type": "Test Case",
        "automation_status": "Not Automated",
        "state": "Design",
        "area_path": "",
        "assigned_to": "",
        "tags": ["Appointments", "Regression", "Smoke"],
    }
    base.update(overrides)
    return GeneratedTestCase(**base)


def _make_entry(test_cases: list[GeneratedTestCase] | None = None, **overrides) -> GenerationHistoryEntry:
    base = {
        "generation_id": "2026-08-03T21-15-42",
        "timestamp": datetime(2026, 8, 3, 21, 15, 42, tzinfo=UTC),
        "feature": "Appointments",
        "redmine_ticket": "12345",
        "model": "gpt-5",
        "prompt_version": "1.0.0",
        "retrieval": RetrievalSummary(
            workflow_chunks=1, historical_test_cases=1, historical_issues=0, uploaded_documents=0
        ),
        "estimated_usage": EstimatedUsage(
            model="gpt-5", input_tokens=1000, estimated_input_cost=MoneyAmount(usd=0.00125, inr=0.11)
        ),
        "actual_usage": ActualUsage(
            model="gpt-5",
            prompt_tokens=1000,
            completion_tokens=200,
            total_tokens=1200,
            input_cost=MoneyAmount(usd=0.00125, inr=0.11),
            output_cost=MoneyAmount(usd=0.002, inr=0.17),
            total_cost=MoneyAmount(usd=0.00325, inr=0.28),
        ),
        "pricing": PricingSnapshot(
            model="gpt-5",
            input_price_per_million_tokens=1.25,
            output_price_per_million_tokens=10.00,
            usd_to_inr_exchange_rate=87.00,
        ),
        "generation_time_ms": 2300.0,
        "prompt": "Prompt Version: 1.0.0\n\n...",
        "response": '{"testCases": []}',
        "test_cases": test_cases if test_cases is not None else [_make_test_case()],
        "metadata": GenerationMetadata(top_k=5, upload_session_id="sess-abc123", generated_test_cases=1),
    }
    base.update(overrides)
    return GenerationHistoryEntry(**base)


def _load_workbook(workbook_bytes: bytes):
    return openpyxl.load_workbook(io.BytesIO(workbook_bytes))


def _rows(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_build_workbook_returns_loadable_xlsx_bytes():
    service = ExcelExportService()

    workbook_bytes = service.build_workbook(_make_entry())

    assert isinstance(workbook_bytes, bytes)
    workbook = _load_workbook(workbook_bytes)
    assert workbook is not None


def test_build_workbook_creates_the_two_expected_sheets():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))

    assert workbook.sheetnames == ["Test Cases", "Generation Summary"]


def test_build_workbook_test_cases_sheet_has_the_exact_header_row():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    header = [cell.value for cell in sheet[1]]
    assert header == _TEST_CASE_COLUMNS


def test_build_workbook_header_row_is_bold():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    assert all(cell.font.bold for cell in sheet[1])


def test_build_workbook_freezes_the_first_row():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    assert sheet.freeze_panes == "A2"


def test_build_workbook_gives_one_test_case_header_row_plus_one_row_per_step():
    test_case = _make_test_case(
        steps=[
            _make_step(step_no=1, action="Step 1 action.", expected_result="Step 1 result."),
            _make_step(step_no=2, action="Step 2 action.", expected_result="Step 2 result."),
            _make_step(step_no=3, action="Step 3 action.", expected_result="Step 3 result."),
        ]
    )
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    data_rows = _rows(sheet)[1:]
    assert len(data_rows) == 4  # 1 header row + 3 step rows
    header_row, *step_rows = data_rows
    assert header_row[2] == "Reschedule an appointment"  # Title
    assert [row[3] for row in step_rows] == [1, 2, 3]  # Test Step
    assert [row[4] for row in step_rows] == ["Step 1 action.", "Step 2 action.", "Step 3 action."]
    assert [row[5] for row in step_rows] == ["Step 1 result.", "Step 2 result.", "Step 3 result."]


def test_build_workbook_leaves_step_columns_blank_on_the_test_case_header_row():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[3] is None  # Test Step
    assert header_row[4] is None  # Step Action
    assert header_row[5] is None  # Step Expected


def test_build_workbook_leaves_test_case_level_columns_blank_on_step_rows():
    test_case = _make_test_case(steps=[_make_step(), _make_step(step_no=2)])
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    data_rows = _rows(sheet)[1:]
    [step_row_1, step_row_2] = data_rows[1:]
    for step_row in (step_row_1, step_row_2):
        assert step_row[0] is None  # ID
        assert step_row[1] is None  # Work Item Type
        assert step_row[2] is None  # Title
        assert step_row[6] is None  # Area Path
        assert step_row[7] is None  # Assigned To
        assert step_row[8] is None  # State
        assert step_row[9] is None  # Tags
        assert step_row[10] is None  # Automation Status


def test_build_workbook_never_fabricates_an_ado_id():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[0] is None  # ID — left for Azure DevOps to assign.


def test_build_workbook_sets_work_item_type_to_test_case():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[1] == "Test Case"


def test_build_workbook_exports_automation_status_as_exactly_not_automated():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[10] == "Not Automated"
    assert header_row[10] not in {"No", "NOT AUTOMATED", "Not automated", "Automated", "Yes"}


def test_build_workbook_includes_every_test_case():
    test_cases = [
        _make_test_case(test_case_id="TC-1", steps=[_make_step()]),
        _make_test_case(test_case_id="TC-2", steps=[_make_step(), _make_step(step_no=2)]),
    ]
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=test_cases)))
    sheet = workbook["Test Cases"]

    data_rows = _rows(sheet)[1:]
    # TC-1: 1 header + 1 step = 2 rows. TC-2: 1 header + 2 steps = 3 rows.
    assert len(data_rows) == 5
    titles = [row[2] for row in data_rows if row[2]]
    assert titles == ["Reschedule an appointment", "Reschedule an appointment"]


def test_build_workbook_gives_a_test_case_with_no_steps_exactly_one_row():
    test_case = _make_test_case(test_case_id="TC-empty", steps=[])
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    data_rows = _rows(sheet)[1:]
    assert len(data_rows) == 1
    assert data_rows[0][2] == "Reschedule an appointment"
    assert data_rows[0][3] is None  # Test Step


def test_build_workbook_joins_multiple_tags_with_semicolon_space():
    test_case = _make_test_case(tags=["Appointments", "Smoke", "Extra"])
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[9] == "Appointments; Smoke; Extra"


def test_build_workbook_formats_a_smoke_test_cases_tags_as_feature_regression_smoke():
    """Regression is the complete suite; Smoke is a subset of it — the
    exported Tags cell for a Smoke test case must never read
    "Feature; Smoke" with no "Regression".
    """
    test_case = _make_test_case(tags=["Contact and Sticket Log", "Regression", "Smoke"])
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[9] == "Contact and Sticket Log; Regression; Smoke"
    assert header_row[9] != "Contact and Sticket Log; Smoke"


def test_build_workbook_formats_a_regression_only_test_cases_tags_as_feature_regression():
    test_case = _make_test_case(tags=["Contact and Sticket Log", "Regression"])
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    [header_row] = _rows(sheet)[1:2]
    assert header_row[9] == "Contact and Sticket Log; Regression"


def test_build_workbook_handles_empty_or_missing_tags():
    for tags in ([], None):
        test_case = _make_test_case(tags=tags)
        service = ExcelExportService()

        workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
        sheet = workbook["Test Cases"]

        [header_row] = _rows(sheet)[1:2]
        assert header_row[9] is None  # openpyxl round-trips an empty-string cell as None


def test_build_workbook_wraps_text_for_long_form_columns():
    service = ExcelExportService()

    test_case = _make_test_case(steps=[_make_step()])
    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    wrap_columns = {"Title": 3, "Step Action": 5, "Step Expected": 6}
    for column_index in wrap_columns.values():
        cell = sheet.cell(row=2, column=column_index)
        assert cell.alignment.wrap_text is True


def test_build_workbook_does_not_wrap_text_for_other_columns():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    cell = sheet.cell(row=2, column=2)  # Work Item Type
    assert not cell.alignment.wrap_text


def test_build_workbook_centers_specific_columns():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    center_columns = {"State": 9, "Automation Status": 11}
    for column_index in center_columns.values():
        cell = sheet.cell(row=2, column=column_index)
        assert cell.alignment.horizontal == "center"


def test_build_workbook_centers_the_test_step_column_on_step_rows():
    test_case = _make_test_case(steps=[_make_step()])
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=[test_case])))
    sheet = workbook["Test Cases"]

    cell = sheet.cell(row=3, column=4)  # Test Step, on the step row
    assert cell.alignment.horizontal == "center"


def test_build_workbook_autosizes_every_column_to_at_least_its_header_width():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Test Cases"]

    for column_index, column_name in enumerate(_TEST_CASE_COLUMNS, start=1):
        letter = sheet.cell(row=1, column=column_index).column_letter
        assert sheet.column_dimensions[letter].width >= len(column_name)


def test_build_workbook_summary_sheet_includes_all_required_fields():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Generation Summary"]

    values = {row[0]: row[1] for row in _rows(sheet)[1:]}
    assert values["Generation ID"] == "2026-08-03T21-15-42"
    assert values["Timestamp"] == "2026-08-03T21:15:42+00:00"
    assert values["Feature"] == "Appointments"
    assert values["Redmine Ticket"] == "12345"
    assert values["Prompt Version"] == "1.0.0"
    assert values["LLM Model"] == "gpt-5"
    assert values["Generation Time (ms)"] == 2300.0
    assert values["Input Tokens"] == 1000
    assert values["Output Tokens"] == 200
    assert values["Total Tokens"] == 1200
    assert values["Input Cost (USD)"] == 0.00125
    assert values["Output Cost (USD)"] == 0.002
    assert values["Total Cost (USD)"] == 0.00325
    assert values["Input Cost (INR)"] == 0.11
    assert values["Output Cost (INR)"] == 0.17
    assert values["Total Cost (INR)"] == 0.28
    assert values["Exchange Rate"] == 87.00
    assert values["Number of Test Cases"] == 1
    assert values["Total Test Steps"] == 1
    assert values["Upload Session"] == "sess-abc123"


def test_build_workbook_summary_sheet_renders_missing_upload_session_as_blank():
    entry = _make_entry(metadata=GenerationMetadata(top_k=None, upload_session_id=None, generated_test_cases=1))
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(entry))
    sheet = workbook["Generation Summary"]

    values = {row[0]: row[1] for row in _rows(sheet)[1:]}
    assert values["Upload Session"] is None  # openpyxl round-trips an empty-string cell as None


def test_build_workbook_summary_sheet_header_row_is_bold():
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry()))
    sheet = workbook["Generation Summary"]

    assert all(cell.font.bold for cell in sheet[1])


def test_build_workbook_counts_total_test_steps_across_all_test_cases():
    test_cases = [
        _make_test_case(test_case_id="TC-1", steps=[_make_step()]),
        _make_test_case(test_case_id="TC-2", steps=[_make_step(), _make_step(step_no=2), _make_step(step_no=3)]),
    ]
    service = ExcelExportService()

    workbook = _load_workbook(service.build_workbook(_make_entry(test_cases=test_cases)))
    sheet = workbook["Generation Summary"]

    values = {row[0]: row[1] for row in _rows(sheet)[1:]}
    assert values["Number of Test Cases"] == 2
    assert values["Total Test Steps"] == 4


def test_build_workbook_raises_conflict_error_when_there_are_no_test_cases():
    service = ExcelExportService()

    with pytest.raises(ConflictError):
        service.build_workbook(_make_entry(test_cases=[]))
